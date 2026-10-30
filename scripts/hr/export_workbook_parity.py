"""Export each employee's monthly gross and deductions from the salary workbook
as JSON, for checking the payroll engine against what was actually deducted.

Usage:
    python scripts/hr/export_workbook_parity.py <workbook.xlsx> <out.json>
    HR_PARITY_JSON=<out.json> npx vitest run --coverage=false src/tests/unit/payrollWorkbookParity.test.ts

The output holds pay figures (no names or ID numbers). Write it outside the repo
and delete it when done.
"""

import json
import sys

# openpyxl imports numpy when present; a broken numpy install must not stop the export.
sys.modules.setdefault("numpy", None)
import openpyxl  # noqa: E402

SKIP_SHEETS = {"SUMMARY", "PAYE PAYMENTS", "SHEET1"}
MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]


def num(v):
    return float(v) if isinstance(v, (int, float)) else None


def main():
    workbook, out = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(workbook, data_only=True)
    employees = []
    index = 0
    for ws in wb.worksheets:
        if ws.title.strip().upper() in SKIP_SHEETS:
            continue
        index += 1
        months = []
        for row in range(11, 40):
            label = ws.cell(row=row, column=1).value
            if not isinstance(label, str) or label.strip() not in MONTHS:
                continue
            gross = num(ws.cell(row=row, column=2).value)
            if gross is None:
                continue
            # Deductions are stored as negatives.
            nis, paye, rr = (-(num(ws.cell(row=row, column=c).value) or 0) for c in (4, 5, 6))
            months.append({"month": MONTHS.index(label.strip()) + 1, "gross": gross, "nis": nis, "paye": paye, "rr": rr})
        employees.append({"employee_no": f"CV-{index:02d}", "annual_allowances": num(ws["K5"].value) or 0, "months": months})

    with open(out, "w", encoding="utf-8") as fh:
        json.dump({"employees": employees}, fh, indent=1)
    print(f"wrote {len(employees)} employees to {out}")


if __name__ == "__main__":
    main()

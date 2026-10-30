"""Turn the yearly salary workbook into SQL that seeds the HR tables.

Usage:
    python scripts/hr/import_salary_xlsx.py <workbook.xlsx> <out.sql> [--year 2026]

The output contains personal data (names, NRN, NIS, pay). Write it outside the
repo, apply it through the Lovable MCP query_database tool, then delete it.
Re-running is safe: employees are keyed by employee_no (CV-<sheet>) and skipped
when they already exist.

Sheets are hand-laid-out and differ per employee, so values are found by their
labels rather than fixed cells. Anything that needs judgement (the PAYMENTS
columns, loan figures) is written to an HR-only note for review instead of
being guessed into pay items.
"""

import argparse
import datetime as dt
import re
import sys

# openpyxl imports numpy when present; a broken numpy install must not stop the import.
sys.modules.setdefault("numpy", None)
import openpyxl  # noqa: E402

SKIP_SHEETS = {"SUMMARY", "PAYE PAYMENTS", "SHEET1"}
NEXT_LABEL = re.compile(r":\s*$|Allowance|Loan|Date Started|Tax Code|^Car$|Benefits|^Other|GROSS|N\.I\.S|N\.R\.N|TAMIS", re.I)


def sql(value):
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return repr(round(float(value), 2))
    if isinstance(value, (dt.date, dt.datetime)):
        return f"'{value:%Y-%m-%d}'"
    return "'" + str(value).replace("'", "''") + "'"


def find_label(ws, pattern, max_row=10):
    rx = re.compile(pattern, re.I)
    for row in ws.iter_rows(min_row=1, max_row=max_row):
        for cell in row:
            if isinstance(cell.value, str) and rx.search(cell.value):
                return cell
    return None


def value_after(ws, pattern, numeric=False, max_cols=8):
    """First non-empty cell to the right of a label (on the same row)."""
    label = find_label(ws, pattern)
    if not label:
        return None
    for col in range(label.column + 1, label.column + 1 + max_cols):
        v = ws.cell(row=label.row, column=col).value
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        if isinstance(v, str) and NEXT_LABEL.search(v):
            return None  # ran into the next label
        if numeric and not isinstance(v, (int, float)):
            continue
        return v
    return None


def split_name(full):
    parts = str(full).split()
    return parts[0].title(), parts[-1].title()


def payment_columns(ws):
    """Labels under PAYMENTS (row 11) with the first month's amounts (row 12)."""
    header = find_label(ws, r"^PAYMENTS$", max_row=10)
    if not header:
        return []
    out = []
    for col in range(header.column, header.column + 4):
        label = ws.cell(row=11, column=col).value
        amount = ws.cell(row=12, column=col).value
        if label is not None or amount is not None:
            out.append((label, amount))
    return out


def employee_sql(ws, index, year):
    full_name = value_after(ws, r"^Name:")
    if not full_name:
        return None, None
    first, last = split_name(full_name)
    sheet_first = ws.title.split()[0].title()
    preferred = sheet_first if sheet_first != first else None
    employee_no = f"CV-{index:02d}"

    base = value_after(ws, r"^Salary Base", numeric=True)
    hire = value_after(ws, r"Date Started")
    tax_code = value_after(ws, r"Tax Code")
    personal = value_after(ws, r"Annual Allowance", numeric=True) or 0
    child = value_after(ws, r"Child Allowance", numeric=True) or 0
    other = value_after(ws, r"Other Allowance", numeric=True) or 0
    car = value_after(ws, r"^Car$", numeric=True)
    loan = value_after(ws, r"Loan Payments", numeric=True)
    nrn = value_after(ws, r"N\.R\.N")
    nis = value_after(ws, r"N\.I\.S")
    tamis = value_after(ws, r"TAMIS")

    note_lines = [f"Imported from the {year} salary workbook (sheet '{ws.title.strip()}')."]
    for label, amount in payment_columns(ws):
        note_lines.append(f"Payments column: {label if label is not None else '(no label)'} = {amount if amount is not None else '-'}")
    if loan:
        note_lines.append(f"Loan payments figure on sheet: {loan}")
    note_lines.append("Set up pay items and net pay distribution from these, then delete this note.")

    stmts = [
        f"""INSERT INTO public.employees (employee_no, first_name, last_name, preferred_name, hire_date, tax_code, status)
SELECT {sql(employee_no)}, {sql(first)}, {sql(last)}, {sql(preferred)}, {sql(hire if isinstance(hire, dt.datetime) else None)}, {sql(None if tax_code is None else str(tax_code))}, 'active'
WHERE NOT EXISTS (SELECT 1 FROM public.employees WHERE employee_no = {sql(employee_no)});""",
        f"""WITH e AS (SELECT id FROM public.employees WHERE employee_no = {sql(employee_no)})
INSERT INTO public.employee_compensation (employee_id, base_amount, annual_personal_allowance, annual_child_allowance, annual_other_allowance, effective_from, reason)
SELECT e.id, {sql(base or 0)}, {sql(personal)}, {sql(child)}, {sql(other)}, '{year}-01-01', 'Imported from salary workbook'
FROM e WHERE NOT EXISTS (SELECT 1 FROM public.employee_compensation c WHERE c.employee_id = e.id);""",
        f"""WITH e AS (SELECT id FROM public.employees WHERE employee_no = {sql(employee_no)})
INSERT INTO public.employee_sensitive (employee_id, national_registration_no, nis_no, tax_file_no)
SELECT e.id, {sql(None if nrn is None else str(nrn))}, {sql(None if nis is None else str(nis))}, {sql(None if tamis is None else str(tamis))}
FROM e ON CONFLICT (employee_id) DO NOTHING;""",
        f"""WITH e AS (SELECT id FROM public.employees WHERE employee_no = {sql(employee_no)})
INSERT INTO public.employee_notes (employee_id, category, body, visibility)
SELECT e.id, 'general', {sql(chr(10).join(note_lines))}, 'hr_only'
FROM e WHERE NOT EXISTS (SELECT 1 FROM public.employee_notes n WHERE n.employee_id = e.id AND n.body LIKE 'Imported from the%');""",
    ]
    if car:
        stmts.append(
            f"""WITH e AS (SELECT id FROM public.employees WHERE employee_no = {sql(employee_no)})
INSERT INTO public.employee_recurring_items (employee_id, kind, label, amount, taxable)
SELECT e.id, 'earning', 'Car allowance', {sql(car)}, true
FROM e WHERE NOT EXISTS (SELECT 1 FROM public.employee_recurring_items r WHERE r.employee_id = e.id AND r.label = 'Car allowance');"""
        )

    summary = f"{employee_no}  base={base}  allowances={personal + child + other}  car={car or 0}  ids={'nrn ' if nrn else ''}{'nis ' if nis else ''}{'tamis' if tamis else ''}"
    return "\n\n".join(stmts), summary


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("out")
    parser.add_argument("--year", type=int, default=dt.date.today().year)
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    chunks = ["BEGIN;"]
    index = 0
    for ws in wb.worksheets:
        if ws.title.strip().upper() in SKIP_SHEETS:
            continue
        index += 1
        stmt, summary = employee_sql(ws, index, args.year)
        if stmt is None:
            print(f"skipped sheet #{index}: no Name: label", file=sys.stderr)
            continue
        chunks.append(f"-- sheet #{index}\n{stmt}")
        # Summary avoids names and ID numbers so it can be pasted into a chat or log.
        print(summary)
    chunks.append("COMMIT;")

    with open(args.out, "w", encoding="utf-8") as fh:
        fh.write("\n\n".join(chunks) + "\n")
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
